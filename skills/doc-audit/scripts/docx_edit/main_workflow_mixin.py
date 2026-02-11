"""
This mixin class implements the overall apply/save flow and omment insertion, manual-action handling, per-item processing.
Handling how a single edit item is processed end-to-end.
"""

from drawing_image_extractor import normalize_drawing_placeholders_in_text

from .common import (
    NS,
    EditItem,
    EditResult,
    COMMENTS_CONTENT_TYPE,
    DEBUG_MARKER,
    filter_synthetic_runs,
)
from typing import List, Dict, Optional, Tuple

from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

from utils import sanitize_xml_string


class MainWorkflowMixin:
    _COMMENT_SOFT_BREAK_TOKEN = "<<<DOC_AUDIT_SOFT_BREAK>>>"
    _STATUS_REASON_SUMMARY_MAX_LEN = 48

    @staticmethod
    def _pick_first_non_none(*candidates):
        """Return the first candidate that is not None."""
        for candidate in candidates:
            if candidate is not None:
                return candidate
        return None

    def _reset_status_reason(self) -> None:
        """Clear pending fallback/conflict reason for current item."""
        self._pending_status_reason = None

    def _summarize_status_reason_detail(self, detail: str) -> str:
        """Build concise single-line summary for long reason detail."""
        text = " ".join((detail or "").split())
        if not text:
            return ""
        for sep in ('; ', '. '):
            if sep in text:
                text = text.split(sep, 1)[0].strip()
                break
        if len(text) > self._STATUS_REASON_SUMMARY_MAX_LEN:
            text = text[: self._STATUS_REASON_SUMMARY_MAX_LEN - 3].rstrip() + "..."
        return text

    def _set_status_reason(self, status: str, code: str, detail: str = "") -> str:
        """Store one-shot structured reason for the next result build."""
        code_text = (code or "").strip() or "FB_UNKNOWN"
        detail_text = " ".join((detail or "").split())
        summary_text = self._summarize_status_reason_detail(detail_text)
        self._pending_status_reason = {
            'status': status,
            'code': code_text,
            'summary': summary_text,
            'detail': detail_text,
        }
        return status

    def _consume_status_reason(self, status: str, include_detail: bool = False) -> Optional[str]:
        """Return and clear pending reason only when status matches."""
        pending = getattr(self, '_pending_status_reason', None)
        if not pending:
            return None
        if pending.get('status') != status:
            return None
        short_code = pending.get('code')
        body = pending.get('detail') if include_detail else pending.get('summary')
        reason = f"{short_code}: {body}" if body else short_code
        self._pending_status_reason = None
        return reason

    def _append_reference_only_comment(
        self,
        para_elem,
        comment_text: str,
        author: str,
        fallback_reason: Optional[str] = None
    ) -> bool:
        """Append a reference-only comment at paragraph end.

        When fallback_reason is provided, the comment text is prefixed with
        [FALLBACK]{reason} as a visible degradation marker.
        """
        if para_elem is None:
            return False

        comment_id = self.next_comment_id
        self.next_comment_id += 1

        ref_xml = f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>'''
        para_elem.append(etree.fromstring(ref_xml))

        full_text = comment_text or ""
        if fallback_reason:
            if full_text:
                full_text = f"[FALLBACK]{fallback_reason}\n{full_text}"
            else:
                full_text = f"[FALLBACK]{fallback_reason}"

        self.comments.append({
            'id': comment_id,
            'text': full_text,
            'author': author
        })
        return True

    def _comment_doc_key_for_run(
        self,
        run_info: Dict,
        default_para
    ) -> Optional[Tuple[int, int]]:
        """Compute document-order key using host para if available, else para."""
        host_para = run_info.get('host_para_elem')
        para = run_info.get('para_elem', default_para)
        key = None
        if host_para is not None:
            key = self._get_run_doc_key(run_info.get('elem'), host_para)
        if key is None and para is not None and para is not host_para:
            key = self._get_run_doc_key(run_info.get('elem'), para)
        return key

    def _pick_reference_para(
        self,
        runs: List[Dict],
        default_para,
        end_host_para=None,
        start_host_para=None
    ):
        """Choose a reference paragraph that has visible text if possible."""
        reference_para = None
        for run_info in reversed(runs):
            para_candidate = run_info.get('host_para_elem')
            if para_candidate is None:
                para_candidate = run_info.get('para_elem')
            if para_candidate is None:
                continue
            try:
                _, para_text = self._collect_runs_info_original(para_candidate)
                if para_text.strip():
                    reference_para = para_candidate
                    break
            except Exception:
                continue

        if reference_para is None:
            reference_para = end_host_para if end_host_para is not None else start_host_para
            if reference_para is None:
                reference_para = default_para
        return reference_para

    def _append_manual_reference_only_comment(
        self,
        para_elem,
        comment_id: int,
        violation_reason: str,
        revised_text: str,
        author: str,
        fallback_reason: str = "FB_REF_ONLY: reference-only comment",
    ) -> bool:
        """Append a reference-only fallback comment using existing comment_id."""
        if para_elem is None:
            print("  [Warning] Reference-only fallback failed: no anchor paragraph")
            return False

        ref_xml = f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>'''
        para_elem.append(etree.fromstring(ref_xml))

        comment_text = f"[FALLBACK]{fallback_reason}\n{violation_reason}"
        if revised_text:
            comment_text += f"\nSuggestion: {revised_text}"

        self.comments.append({
            'id': comment_id,
            'text': comment_text,
            'author': author
        })
        return True

    def _build_manual_comment_text(
        self,
        violation_reason: str,
        revised_text: str,
        fallback_reason: Optional[str] = None
    ) -> str:
        """Build final comment text for manual action."""
        if fallback_reason:
            comment_text = f"[FALLBACK]{fallback_reason}\n{violation_reason}"
        else:
            comment_text = violation_reason
        if revised_text:
            comment_text += f"\nSuggestion: {revised_text}"
        return comment_text

    def _insert_manual_start_marker(
        self,
        range_start,
        start_use_host: bool,
        start_host_para,
        start_revision,
        first_run,
        no_split_start: bool,
        before_text: str,
        rPr_xml: str
    ) -> bool:
        """Insert commentRangeStart marker for manual comment."""
        if start_use_host:
            target_para = start_host_para
            if target_para is None:
                print("  [Warning] Reference-only fallback failed: no anchor paragraph")
                return False
            insert_idx = 0
            for i, child in enumerate(list(target_para)):
                if child.tag == f'{{{NS["w"]}}}pPr':
                    continue
                insert_idx = i
                break
            target_para.insert(insert_idx, range_start)
            return True

        if start_revision is not None:
            parent = start_revision.getparent()
            if parent is not None:
                idx = list(parent).index(start_revision)
                parent.insert(idx, range_start)
            return True

        if first_run is None:
            return False

        parent = first_run.getparent()
        if parent is None:
            return False
        try:
            idx = list(parent).index(first_run)
        except ValueError:
            return False

        if no_split_start:
            parent.insert(idx, range_start)
            return True

        if before_text:
            run_or_container = self._create_run(before_text, rPr_xml)
            if run_or_container.tag == 'container':
                for child_run in run_or_container:
                    parent.insert(idx, child_run)
                    idx += 1
            else:
                parent.insert(idx, run_or_container)
                idx += 1

            t_elem = first_run.find('w:t', NS)
            if t_elem is not None and t_elem.text:
                t_elem.text = t_elem.text[len(before_text):]

        parent.insert(idx, range_start)
        return True

    def _insert_manual_end_marker(
        self,
        range_end,
        comment_ref,
        end_use_host: bool,
        reference_para,
        end_host_para,
        end_revision,
        last_run,
        after_text: str,
        rPr_xml: str,
        no_split_end: bool
    ) -> bool:
        """Insert commentRangeEnd and commentReference for manual comment."""
        if end_use_host:
            target_para = reference_para if reference_para is not None else end_host_para
            if target_para is None:
                print("  [Warning] Reference-only fallback failed: no anchor paragraph")
                return False
            if (
                self.verbose and reference_para is not None and end_host_para is not None
                and reference_para is not end_host_para
            ):
                print("  [Debug] End anchor moved to previous non-empty paragraph")
            target_para.append(range_end)
            target_para.append(comment_ref)
            return True

        if end_revision is not None:
            parent = end_revision.getparent()
            if parent is not None:
                idx = list(parent).index(end_revision)
                parent.insert(idx + 1, range_end)
                parent.insert(idx + 2, comment_ref)
            return True

        parent = last_run.getparent() if last_run is not None else None
        if parent is None:
            return False
        try:
            idx = list(parent).index(last_run)
        except ValueError:
            return False

        if no_split_end:
            parent.insert(idx + 1, range_end)
            parent.insert(idx + 2, comment_ref)
            return True

        if after_text:
            t_elem = last_run.find('w:t', NS)
            if t_elem is not None and t_elem.text:
                original_text = t_elem.text
                keep_len = len(original_text) - len(after_text)
                t_elem.text = original_text[:keep_len]

            parent.insert(idx + 1, range_end)
            parent.insert(idx + 2, comment_ref)
            run_or_container = self._create_run(after_text, rPr_xml)
            if run_or_container.tag == 'container':
                insert_pos = idx + 3
                for child_run in run_or_container:
                    parent.insert(insert_pos, child_run)
                    insert_pos += 1
            else:
                parent.insert(idx + 3, run_or_container)
            return True

        parent.insert(idx + 1, range_end)
        parent.insert(idx + 2, comment_ref)
        return True

    def _insert_manual_style_range_comment(
        self,
        real_runs: List[Dict],
        para_elem,
        violation_reason: str,
        author: str
    ) -> Tuple[bool, str]:
        """Insert range comment markers using manual-comment style anchoring."""
        if not real_runs:
            return False, "No real runs available"

        first_run_info = real_runs[0]
        last_run_info = real_runs[-1]

        start_key = self._comment_doc_key_for_run(first_run_info, para_elem)
        end_key = self._comment_doc_key_for_run(last_run_info, para_elem)
        if start_key is None:
            return False, "Cannot resolve start run order"

        if end_key is None or (end_key is not None and end_key < start_key):
            candidate = None
            for run_info in reversed(real_runs):
                cand_key = self._comment_doc_key_for_run(run_info, para_elem)
                if cand_key is not None and cand_key >= start_key:
                    candidate = run_info
                    break
            if candidate is None:
                return False, "Cannot resolve valid end run order"
            last_run_info = candidate

        first_run = first_run_info.get('elem')
        last_run = last_run_info.get('elem')

        start_host_para = first_run_info.get('host_para_elem')
        end_host_para = last_run_info.get('host_para_elem')
        start_use_host = start_host_para is not None and start_host_para is not first_run_info.get('para_elem')
        end_use_host = end_host_para is not None and end_host_para is not last_run_info.get('para_elem')

        has_host_mismatch = any(
            r.get('host_para_elem') is not None and r.get('host_para_elem') is not r.get('para_elem')
            for r in real_runs
        )
        if has_host_mismatch:
            start_use_host = True
            end_use_host = True

        if first_run is None:
            start_use_host = True
        if last_run is None:
            end_use_host = True

        if start_use_host and start_host_para is None:
            start_host_para = first_run_info.get('para_elem', para_elem)
        if end_use_host and end_host_para is None:
            end_host_para = last_run_info.get('para_elem', para_elem)

        reference_para = self._pick_reference_para(
            real_runs,
            para_elem,
            end_host_para=end_host_para,
            start_host_para=start_host_para
        )

        if start_use_host and end_use_host and start_host_para is not None and reference_para is not None:
            self._init_para_order()
            start_idx = self._para_order.get(id(start_host_para))
            end_idx = self._para_order.get(id(reference_para))
            if start_idx is not None and end_idx is not None and end_idx < start_idx:
                return False, "Host paragraph order is inverted"

        if start_use_host:
            target_para = start_host_para
            if target_para is None:
                return False, "Missing start host paragraph"
            start_parent = target_para
            start_index = 0
            for i, child in enumerate(list(target_para)):
                if child.tag == f'{{{NS["w"]}}}pPr':
                    continue
                start_index = i
                break
        else:
            first_para = first_run_info.get('para_elem', para_elem)
            start_revision = self._find_revision_ancestor(first_run, first_para)
            if start_revision is not None:
                start_parent = start_revision.getparent()
                if start_parent is None:
                    return False, "Cannot place start marker outside revision"
                start_index = list(start_parent).index(start_revision)
            else:
                start_parent = first_run.getparent() if first_run is not None else None
                if start_parent is None:
                    return False, "Missing start run parent"
                start_index = list(start_parent).index(first_run)

        if end_use_host:
            target_para = reference_para if reference_para is not None else end_host_para
            if target_para is None:
                return False, "Missing end host paragraph"
            end_parent = target_para
            end_index = len(list(target_para))
        else:
            last_para = last_run_info.get('para_elem', para_elem)
            end_revision = self._find_revision_ancestor(last_run, last_para)
            if end_revision is not None:
                end_parent = end_revision.getparent()
                if end_parent is None:
                    return False, "Cannot place end marker outside revision"
                end_index = list(end_parent).index(end_revision) + 1
            else:
                end_parent = last_run.getparent() if last_run is not None else None
                if end_parent is None:
                    return False, "Missing end run parent"
                end_index = list(end_parent).index(last_run) + 1

        comment_id = self.next_comment_id
        self.next_comment_id += 1
        range_start = etree.fromstring(
            f'<w:commentRangeStart xmlns:w="{NS["w"]}" w:id="{comment_id}"/>'
        )
        range_end = etree.fromstring(
            f'<w:commentRangeEnd xmlns:w="{NS["w"]}" w:id="{comment_id}"/>'
        )
        comment_ref = etree.fromstring(f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>''')

        inserted = []
        try:
            start_parent.insert(start_index, range_start)
            inserted.append(range_start)

            if start_parent is end_parent and start_index < end_index:
                end_index += 1

            end_parent.insert(end_index, range_end)
            inserted.append(range_end)
            end_parent.insert(end_index + 1, comment_ref)
            inserted.append(comment_ref)
        except Exception:
            for node in reversed(inserted):
                parent = node.getparent()
                if parent is not None:
                    parent.remove(node)
            self.next_comment_id = comment_id
            return False, "Failed to insert range markers"

        self.comments.append({
            'id': comment_id,
            'text': violation_reason,
            'author': author
        })
        return True, ""

    def _try_add_multi_cell_range_comment(
        self,
        success_cell_edits: List[Dict],
        violation_reason: str,
        author: str
    ) -> Tuple[bool, str]:
        """
        Try inserting a range comment for multi-cell replace.

        Strategy:
        - Re-collect fresh runs from successful cells after replace.
        - Build first/last anchors in document order.
        - Insert markers using the same anchoring rules as manual comments.

        Returns:
            (True, "") on success, (False, reason) on fallback.
        """
        if not success_cell_edits:
            return False, "No successful cells available"

        cells = []
        seen_cells = set()
        for cell_edit in success_cell_edits:
            cell_elem = cell_edit.get('cell_elem')
            if cell_elem is None:
                return False, "Missing cell element"
            cell_id = id(cell_elem)
            if cell_id in seen_cells:
                continue
            seen_cells.add(cell_id)
            cells.append(cell_elem)

        if not cells:
            return False, "No valid cells for range comment"

        candidate_runs = []
        seen_run_keys = set()
        for cell_elem in cells:
            paras = self._xpath(cell_elem, './/w:p')
            for para in paras:
                para_runs, _ = self._collect_runs_info_original(para)
                for run in para_runs:
                    run_elem = run.get('elem')
                    if run_elem is None:
                        continue
                    run_key = (id(para), id(run_elem))
                    if run_key in seen_run_keys:
                        continue
                    seen_run_keys.add(run_key)

                    run['para_elem'] = para
                    run['host_para_elem'] = para
                    run['cell_elem'] = cell_elem
                    run['row_elem'] = self._find_ancestor_row(para)
                    candidate_runs.append(run)

        if not candidate_runs:
            return False, "No run anchor found in successful cells"

        keyed_runs = []
        for run in candidate_runs:
            para = run.get('host_para_elem')
            if para is None:
                para = run.get('para_elem')
            key = self._get_run_doc_key(run.get('elem'), para)
            if key is not None:
                keyed_runs.append((key, run))

        if not keyed_runs:
            return False, "Cannot resolve run order for successful cells"

        keyed_runs.sort(key=lambda item: item[0])
        ordered_runs = [run for _, run in keyed_runs]
        anchor_para = ordered_runs[0].get('para_elem')
        return self._insert_manual_style_range_comment(
            ordered_runs, anchor_para, violation_reason, author
        )

    def _apply_error_comment(self, para_elem, item: EditItem, author_override: str = None) -> bool:
        """
        Insert an unselected comment at the end of paragraph for failed items.
        
        Comment format: {WHY}<violation_reason>  {WHERE}<violation_text>{SUGGEST}<revised_text>
        Author: author_override if provided, otherwise {self.author}-{category}
        
        Args:
            para_elem: Paragraph element to attach comment
            item: EditItem with violation details
            author_override: Optional author name to use instead of default
        """
        comment_id = self.next_comment_id
        self.next_comment_id += 1
        
        # Insert only commentReference at the end of paragraph (no range selection)
        ref_xml = f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>'''
        para_elem.append(etree.fromstring(ref_xml))
        
        # Record comment content with custom format and author
        comment_text = f"{{WHY}}{item.violation_reason}  {{WHERE}}{item.violation_text}{{SUGGEST}}{item.revised_text}"
        comment_author = author_override if author_override else self._author_for_item(item)
        self.comments.append({
            'id': comment_id,
            'text': comment_text,
            'author': comment_author
        })
        
        return True

    def _apply_fallback_comment(self, para_elem, item: EditItem, reason: str = "") -> bool:
        """
        Insert a fallback comment when delete/replace/manual operation cannot be applied.
        
        This is different from _apply_error_comment:
        - Error comment: Unexpected failure
        - Fallback comment: Expected fallback (text was modified by previous rule)
        
        Args:
            para_elem: Paragraph element
            item: Edit item
            reason: Reason for fallback (e.g., "Text modified by previous rule")
        
        Returns:
            True (always succeeds)
        """
        comment_id = self.next_comment_id
        self.next_comment_id += 1
        
        # Insert commentReference at end of paragraph
        ref_xml = f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>'''
        para_elem.append(etree.fromstring(ref_xml))
        
        # Format: [FALLBACK] reason | {WHY} ... {WHERE} ... {SUGGEST} ...
        comment_text = f"[FALLBACK]{reason}\n{{WHY}}{item.violation_reason} {{WHERE}}{item.violation_text} {{SUGGEST}}{item.revised_text}"
        
        self.comments.append({
            'id': comment_id,
            'text': comment_text,
            'author': self._author_for_item(item)
        })
        
        return True

    def _apply_cell_fallback_comment(self, para_elem, cell_violation: str, 
                                      cell_revised: str, reason: str, 
                                      item: EditItem) -> bool:
        """
        Insert a fallback comment for a single failed cell in multi-cell operation.
        
        Args:
            para_elem: Paragraph element in the failed cell
            cell_violation: Cell-specific violation text
            cell_revised: Cell-specific revised text
            reason: Reason for failure (e.g., "Text not found")
            item: Original edit item (for rule_id and violation_reason)
        
        Returns:
            True (always succeeds)
        """
        comment_id = self.next_comment_id
        self.next_comment_id += 1
        
        # Insert commentReference at end of paragraph
        ref_xml = f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>'''
        para_elem.append(etree.fromstring(ref_xml))
        
        # Format: [CELL FAILED] reason | {WHY} ... {WHERE} ... {SUGGEST} ...
        comment_text = f"[CELL FAILED] {reason}\n{{WHY}}{item.violation_reason}  {{WHERE}}{cell_violation}{{SUGGEST}}{cell_revised}"
        
        self.comments.append({
            'id': comment_id,
            'text': comment_text,
            'author': self._author_for_item(item)
        })
        
        return True

    def _apply_manual(self, para_elem, violation_text: str,
                     violation_reason: str, revised_text: str,
                     orig_runs_info: List[Dict],
                     orig_match_start: int,
                     author: str,
                     is_cross_paragraph: bool = False,
                     fallback_reason: Optional[str] = None) -> str:
        """
        Apply manual operation by adding a Word comment.
        
        This method preserves internal run structure (including images, revision 
        markup, etc.) by inserting comment markers at element boundaries rather
        than replacing the content.
        
        Strategy:
        1. Find the precise start/end positions in the original text
        2. For each boundary (start/end), check if the run is inside revision markup
        3. If inside revision: insert marker outside the revision container
        4. If not inside revision: split run if needed, insert marker at precise position
        
        Cross-paragraph support:
        - When is_cross_paragraph=True, commentRangeStart/End can span multiple paragraphs
        - Uses 'para_elem' field from runs to determine which paragraph to insert markers into
        
        Args:
            para_elem: Paragraph element (anchor paragraph, may not be used in cross-para mode)
            violation_text: Text to mark with comment
            violation_reason: Reason to show in comment
            revised_text: Suggestion to show in comment
            orig_runs_info: Pre-computed original runs info from _process_item
            orig_match_start: Pre-computed match position in original text
            author: Comment author (base author + category suffix)
            is_cross_paragraph: If True, handle cross-paragraph comment range
        
        Returns:
            'success': Comment added successfully
            'fallback': Should fallback to error comment
        """
        match_start = orig_match_start
        match_end = match_start + len(violation_text)
        affected = self._find_affected_runs(orig_runs_info, match_start, match_end)
        if not affected:
            return self._set_status_reason(
                'fallback',
                'FB_MAN_NO_HIT',
                'manual hit not found',
            )

        real_runs = self._filter_real_runs(affected, include_equations=True)
        if not real_runs:
            return self._set_status_reason(
                'fallback',
                'FB_MAN_NO_RUN',
                'manual no editable runs',
            )

        comment_id = self.next_comment_id
        self.next_comment_id += 1

        first_run_info = real_runs[0]
        last_run_info = real_runs[-1]
        start_key = self._comment_doc_key_for_run(first_run_info, para_elem)
        end_key = self._comment_doc_key_for_run(last_run_info, para_elem)

        if start_key is not None and end_key is not None and end_key < start_key:
            candidate = None
            for run_info in reversed(real_runs):
                cand_key = self._comment_doc_key_for_run(run_info, para_elem)
                if cand_key is not None and cand_key >= start_key:
                    candidate = run_info
                    break
            if candidate is None:
                fallback_para = self._pick_first_non_none(
                    first_run_info.get('host_para_elem'),
                    first_run_info.get('para_elem'),
                    para_elem,
                )
                if not self._append_manual_reference_only_comment(
                    fallback_para, comment_id, violation_reason, revised_text, author
                ):
                    return self._set_status_reason(
                        'fallback',
                        'FB_MAN_REF_FAIL',
                        'ref-only append failed',
                    )
                if self.verbose:
                    print("  [Debug] ParaId order inverted; fallback to reference-only comment")
                return 'success'
            if self.verbose:
                print("  [Debug] ParaId order inverted; adjusted end run")
            last_run_info = candidate
            match_end = min(match_end, last_run_info.get('end', match_end))

        first_run = first_run_info.get('elem')
        last_run = last_run_info.get('elem')
        rPr_xml = self._get_rPr_xml(first_run_info.get('rPr'))
        no_split_start = first_run_info.get('is_equation', False)
        no_split_end = last_run_info.get('is_equation', False)

        start_host_para = first_run_info.get('host_para_elem')
        end_host_para = last_run_info.get('host_para_elem')
        start_use_host = (
            start_host_para is not None
            and start_host_para is not first_run_info.get('para_elem')
        )
        end_use_host = (
            end_host_para is not None
            and end_host_para is not last_run_info.get('para_elem')
        )

        has_host_mismatch = any(
            r.get('host_para_elem') is not None
            and r.get('host_para_elem') is not r.get('para_elem')
            for r in real_runs
        )
        if has_host_mismatch:
            start_use_host = True
            end_use_host = True

        if first_run is None:
            start_use_host = True
            no_split_start = True
        if last_run is None:
            end_use_host = True
            no_split_end = True

        if start_use_host and start_host_para is None:
            start_host_para = first_run_info.get('para_elem', para_elem)
        if end_use_host and end_host_para is None:
            end_host_para = last_run_info.get('para_elem', para_elem)

        reference_para = self._pick_reference_para(
            real_runs,
            para_elem,
            end_host_para=end_host_para,
            start_host_para=start_host_para
        )

        if (
            start_use_host and end_use_host and start_host_para is not None
            and reference_para is not None
        ):
            self._init_para_order()
            start_idx = self._para_order.get(id(start_host_para))
            end_idx = self._para_order.get(id(reference_para))
            if start_idx is not None and end_idx is not None and end_idx < start_idx:
                fallback_para = start_host_para if start_host_para is not None else para_elem
                if not self._append_manual_reference_only_comment(
                    fallback_para, comment_id, violation_reason, revised_text, author
                ):
                    return self._set_status_reason(
                        'fallback',
                        'FB_MAN_REF_FAIL',
                        'ref-only append failed',
                    )
                if self.verbose:
                    print("  [Debug] ParaId order inverted; fallback to reference-only comment")
                return 'success'

        if is_cross_paragraph:
            first_para = first_run_info.get('para_elem', para_elem)
            last_para = last_run_info.get('para_elem', para_elem)
        else:
            first_para = para_elem
            last_para = para_elem

        start_revision = self._find_revision_ancestor(first_run, first_para)
        end_revision = self._find_revision_ancestor(last_run, last_para)

        first_orig_text = self._get_run_original_text(first_run_info)
        last_orig_text = self._get_run_original_text(last_run_info)
        before_offset = self._translate_escaped_offset(
            first_run_info, max(0, match_start - first_run_info['start'])
        )
        after_offset = self._translate_escaped_offset(
            last_run_info, max(0, match_end - last_run_info['start'])
        )
        before_text = first_orig_text[:before_offset]
        after_text = last_orig_text[after_offset:]

        range_start = etree.fromstring(
            f'<w:commentRangeStart xmlns:w="{NS["w"]}" w:id="{comment_id}"/>'
        )
        range_end = etree.fromstring(
            f'<w:commentRangeEnd xmlns:w="{NS["w"]}" w:id="{comment_id}"/>'
        )
        comment_ref = etree.fromstring(f'''<w:r xmlns:w="{NS['w']}">
            <w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>
            <w:commentReference w:id="{comment_id}"/>
        </w:r>''')

        if not self._insert_manual_start_marker(
            range_start,
            start_use_host=start_use_host,
            start_host_para=start_host_para,
            start_revision=start_revision,
            first_run=first_run,
            no_split_start=no_split_start,
            before_text=before_text,
            rPr_xml=rPr_xml
        ):
            return self._set_status_reason(
                'fallback',
                'FB_MAN_START_FAIL',
                'start marker insert failed',
            )

        if not self._insert_manual_end_marker(
            range_end,
            comment_ref,
            end_use_host=end_use_host,
            reference_para=reference_para,
            end_host_para=end_host_para,
            end_revision=end_revision,
            last_run=last_run,
            after_text=after_text,
            rPr_xml=rPr_xml,
            no_split_end=no_split_end
        ):
            return self._set_status_reason(
                'fallback',
                'FB_MAN_END_FAIL',
                'end marker insert failed',
            )

        self.comments.append({
            'id': comment_id,
            'text': self._build_manual_comment_text(
                violation_reason,
                revised_text,
                fallback_reason=fallback_reason
            ),
            'author': author
        })
        return 'success'

    def _prepare_comment_text_for_word(self, comment_text: str) -> str:
        """
        Prepare comment text before writing to comments.xml.

        For [FALLBACK] comments, only the first newline is rendered as a real
        soft break (<w:br/>). Remaining newlines are escaped as literal "\\n".
        """
        normalized = (comment_text or "").replace('\r\n', '\n').replace('\r', '\n')
        if not normalized.startswith("[FALLBACK]"):
            return normalized
        if '\n' not in normalized:
            return normalized

        head, tail = normalized.split('\n', 1)
        escaped_tail = tail.replace('\n', '\\n')
        return f"{head}{self._COMMENT_SOFT_BREAK_TOKEN}{escaped_tail}"

    def _append_comment_segments_to_paragraph(self, p_elem, prepared_text: str) -> None:
        """Append parsed comment runs to a comment paragraph, supporting soft breaks."""
        segments = self._parse_formatted_text(prepared_text)
        token = self._COMMENT_SOFT_BREAK_TOKEN

        for segment_text, vert_align in segments:
            if segment_text is None:
                continue

            parts = segment_text.split(token)
            for idx, part in enumerate(parts):
                if part:
                    r = etree.SubElement(p_elem, f'{{{NS["w"]}}}r')

                    if vert_align:
                        rPr = etree.SubElement(r, f'{{{NS["w"]}}}rPr')
                        vert_elem = etree.SubElement(rPr, f'{{{NS["w"]}}}vertAlign')
                        vert_elem.set(f'{{{NS["w"]}}}val', vert_align)

                    t = etree.SubElement(r, f'{{{NS["w"]}}}t')
                    # Preserve whitespace (Word drops leading/trailing spaces without this)
                    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    t.text = part

                if idx < len(parts) - 1:
                    br_run = etree.SubElement(p_elem, f'{{{NS["w"]}}}r')
                    etree.SubElement(br_run, f'{{{NS["w"]}}}br')

    def _save_comments(self):
        """Save comments to comments.xml using OPC API"""
        if not self.comments:
            return

        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        
        # Try to get existing comments.xml
        try:
            comments_part = self.doc.part.part_related_by(RT.COMMENTS)
            comments_xml = etree.fromstring(comments_part.blob)
        except (KeyError, AttributeError):
            # Create new comments.xml
            comments_xml = etree.fromstring(
                f'<w:comments xmlns:w="{NS["w"]}" xmlns:w14="{NS["w14"]}"/>'
            )
        
        # Add comments
        for comment in self.comments:
            comment_elem = etree.SubElement(
                comments_xml, f'{{{NS["w"]}}}comment'
            )
            comment_elem.set(f'{{{NS["w"]}}}id', str(comment['id']))
            # Support independent author for each comment (author-category suffix)
            comment_author = comment.get('author', self.author)
            comment_elem.set(f'{{{NS["w"]}}}author', comment_author)
            comment_elem.set(f'{{{NS["w"]}}}date', self.operation_timestamp)
            # Use self.initials for all comments with author prefix matching self.author
            # This includes: AI-<category> (all share same initials)
            if comment_author.startswith(self.author):
                comment_initials = self.initials
            else:
                comment_initials = comment_author[:2] if len(comment_author) >= 2 else comment_author
            comment_elem.set(f'{{{NS["w"]}}}initials', comment_initials)

            # Add paragraph with formatted text (handle <sup>/<sub> markup)
            p = etree.SubElement(comment_elem, f'{{{NS["w"]}}}p')

            # sanitize_xml_string first to remove illegal control characters
            sanitized_text = sanitize_xml_string(comment.get('text', ''))
            prepared_text = self._prepare_comment_text_for_word(sanitized_text)
            self._append_comment_segments_to_paragraph(p, prepared_text)

        if self.verbose:
            try:
                total = len(comments_xml.findall(f'{{{NS["w"]}}}comment'))
                print(f"[Comments] comments.xml count after merge: {total}")
            except Exception:
                pass
        
        # Save via OPC
        blob = etree.tostring(
            comments_xml, xml_declaration=True, encoding='UTF-8'
        )

        try:
            comments_part = self.doc.part.part_related_by(RT.COMMENTS)
            comments_part._blob = blob
        except (KeyError, AttributeError):
            # Create new Part
            comments_part = Part(
                PackURI('/word/comments.xml'),
                COMMENTS_CONTENT_TYPE,
                blob,
                self.doc.part.package
            )
            self.doc.part.relate_to(comments_part, RT.COMMENTS)

        if self.verbose:
            try:
                print(f"[Comments] Saved {len(self.comments)} comment(s) to comments.xml")
            except Exception:
                pass

    def _process_item(self, item: EditItem) -> EditResult:
        """Process a single edit item"""
        anchor_para = None
        try:
            self._reset_status_reason()
            # Strip leading/trailing whitespace from search and replacement text, normalize drawing attribute
            # to prevent matching failures caused by whitespace in JSONL data
            violation_text = normalize_drawing_placeholders_in_text(
                item.violation_text.strip(),
                include_extended_attrs=False,
            )
            revised_text = normalize_drawing_placeholders_in_text(
                item.revised_text.strip(),
                include_extended_attrs=False,
            )
            
            # 1. Find anchor paragraph by ID
            anchor_para = self._find_para_node_by_id(item.uuid)
            if anchor_para is None:
                return EditResult(False, item,
                    f"Paragraph ID {item.uuid} not found (may be in header/footer or ID changed)")
            
            item_author = self._author_for_item(item)

            search_context = self._locate_item_match(
                item=item,
                anchor_para=anchor_para,
                violation_text=violation_text,
                revised_text=revised_text,
            )
            early_result = search_context['early_result']
            if early_result is not None:
                return early_result

            target_para = search_context['target_para']
            matched_runs_info = search_context['matched_runs_info']
            matched_start = search_context['matched_start']
            violation_text = search_context['violation_text']
            revised_text = search_context['revised_text']
            numbering_stripped = search_context['numbering_stripped']
            is_cross_paragraph = search_context['is_cross_paragraph']
            
            # 3. Apply operation based on fix_action
            # Pass matched_runs_info and matched_start to avoid double matching.
            # For delete/replace: if current match conflicts with existing revisions,
            # retry with the next match in the remaining block content.
            def _resolve_target_para_for_match(current_runs: List[Dict], current_start: int,
                                               fallback_para):
                match_end = current_start + len(violation_text)
                affected = self._find_affected_runs(current_runs, current_start, match_end)
                real_runs = filter_synthetic_runs(affected, include_equations=False)
                for run in real_runs:
                    para = run.get('para_elem')
                    if para is not None:
                        return para
                return fallback_para

            def _apply_delete_or_replace(action: str, current_start: int, current_target_para):
                current_status = None
                current_para = current_target_para

                if action == 'delete':
                    # Check for cross-paragraph: only fallback if ACTUAL match spans multiple paragraphs
                    if is_cross_paragraph:
                        # Get affected runs in the matched range
                        match_end = current_start + len(violation_text)
                        affected = self._find_affected_runs(matched_runs_info, current_start, match_end)

                        # Filter out boundary markers (paragraph and JSON boundaries)
                        real_runs = filter_synthetic_runs(affected, include_equations=False)

                        # Check if actual match spans multiple paragraphs
                        para_elems = set(r.get('para_elem') for r in real_runs if r.get('para_elem') is not None)

                        # Check if match spans multiple table rows
                        if self._check_cross_row_boundary(real_runs):
                            # Multi-row delete: use cell-by-cell deletion
                            if self.verbose:
                                print(f"  [Multi-row] Applying cell-by-cell deletion")
                            current_status = self._apply_delete_multi_cell(
                                real_runs, violation_text,
                                item.violation_reason, item_author, current_start,
                                item
                            )
                        # Check if match spans multiple table cells (within same row)
                        elif self._check_cross_cell_boundary(real_runs):
                            # Multi-cell delete: use cell-by-cell deletion
                            if self.verbose:
                                print(f"  [Multi-cell] Applying cell-by-cell deletion")
                            current_status = self._apply_delete_multi_cell(
                                real_runs, violation_text,
                                item.violation_reason, item_author, current_start,
                                item
                            )
                        elif len(para_elems) > 1:
                            # Actually spans multiple paragraphs
                            if self._is_table_mode(real_runs) or any(r.get('cell_elem') is not None for r in real_runs):
                                if self.verbose:
                                    print(f"  [Cross-paragraph] delete spans {len(para_elems)} paragraphs, fallback to comment")
                                current_status = self._set_status_reason(
                                    'cross_paragraph_fallback',
                                    'CP_TBL_SPAN',
                                    f'delete spans {len(para_elems)} paras in table',
                                )
                            else:
                                if self.verbose:
                                    print(f"  [Cross-paragraph] delete spans {len(para_elems)} paragraphs, applying cross-paragraph delete")
                                current_status = self._apply_delete_cross_paragraph(
                                    matched_runs_info, current_start, violation_text,
                                    item.violation_reason, item_author
                                )
                        else:
                            # All content is in single paragraph - safe to delete
                            if real_runs:
                                current_para = real_runs[0].get('para_elem', current_para)
                            current_status = self._apply_delete(
                                current_para, violation_text,
                                item.violation_reason,
                                matched_runs_info, current_start,
                                item_author
                            )
                    else:
                        current_para = _resolve_target_para_for_match(
                            matched_runs_info, current_start, current_para
                        )
                        current_status = self._apply_delete(
                            current_para, violation_text,
                            item.violation_reason,
                            matched_runs_info, current_start,
                            item_author
                        )
                else:
                    # replace
                    # Check for cross-paragraph: only fallback if ACTUAL match spans multiple paragraphs
                    if is_cross_paragraph:
                        # Get affected runs in the matched range
                        match_end = current_start + len(violation_text)
                        affected = self._find_affected_runs(matched_runs_info, current_start, match_end)

                        # Filter out boundary markers (paragraph and JSON boundaries)
                        real_runs = filter_synthetic_runs(affected, include_equations=False)

                        # Check if actual match spans multiple paragraphs
                        para_elems = set(r.get('para_elem') for r in real_runs if r.get('para_elem') is not None)

                        # Check if match spans multiple table rows
                        is_cross_row = self._check_cross_row_boundary(real_runs)

                        # Detect JSON cell boundaries even if only one cell has text runs.
                        # This handles cases where empty cells (e.g., stripped row numbers)
                        # produce boundary markers but no runs with cell_elem.
                        has_cell_boundary = any(r.get('is_cell_boundary') for r in affected)
                        is_cross_cell = self._check_cross_cell_boundary(real_runs) or has_cell_boundary or is_cross_row

                        if is_cross_row and self.verbose:
                            print(f"  [Cross-row] replace spans multiple rows, trying cell-by-cell extraction...")
                        if self.verbose and has_cell_boundary and not self._check_cross_cell_boundary(real_runs):
                            print(f"  [Cross-cell] Boundary markers detected (empty cells), forcing cell extraction")

                        if current_status is None and is_cross_cell:
                            # Try to extract single-cell edit first
                            if self.verbose:
                                print(f"  [Cross-cell] Detected cross-cell match, trying cell-by-cell extraction...")

                            single_cell = self._try_extract_single_cell_edit(
                                violation_text, revised_text, affected, current_start
                            )

                            if single_cell:
                                # Successfully extracted single-cell edit - all changes in one cell
                                current_status = self._apply_replace_in_cell_paragraphs(
                                    single_cell['cell_violation'],
                                    single_cell['cell_revised'],
                                    single_cell['cell_runs'],
                                    item.violation_reason,
                                    item_author,
                                    skip_comment=False
                                )
                                if self.verbose and current_status == 'success':
                                    print(f"  [Single-cell] All changes in one cell, applied track change")
                                elif current_status not in ('success', 'conflict'):
                                    nested_reason = self._consume_status_reason(current_status)
                                    detail = f'single-cell replace returned {current_status}'
                                    if nested_reason:
                                        detail = f"{detail}; {nested_reason}"
                                    current_status = self._set_status_reason(
                                        'cross_cell_fallback',
                                        'CC_SINGLE_FAIL',
                                        detail,
                                    )
                            else:
                                # Try multi-cell extraction - changes distributed across cells
                                multi_cells = self._try_extract_multi_cell_edits(
                                    violation_text, revised_text, affected, current_start
                                )

                                if multi_cells:
                                    # Successfully extracted multi-cell edits - each change within its own cell
                                    if self.verbose:
                                        print(f"  [Multi-cell] Found {len(multi_cells)} cells with changes, applying track changes...")

                                    success_count = 0
                                    failed_cells = []  # List of (cell_edit, cell_para, error_reason)
                                    success_cells = []
                                    first_success_para = None

                                    for cell_edit in multi_cells:
                                        cell_status = self._apply_replace_in_cell_paragraphs(
                                            cell_edit['cell_violation'],
                                            cell_edit['cell_revised'],
                                            cell_edit['cell_runs'],
                                            item.violation_reason,
                                            item_author,
                                            skip_comment=True
                                        )

                                        if cell_status != 'success':
                                            # Get first paragraph for error tracking
                                            first_para = None
                                            for run in cell_edit['cell_runs']:
                                                para = run.get('para_elem')
                                                if para is not None:
                                                    first_para = para
                                                    break
                                            failed_cells.append((cell_edit, first_para, f"Apply failed: {cell_status}"))
                                            continue  # Continue processing next cell

                                        # Success - track for overall comment
                                        success_count += 1
                                        success_cells.append(cell_edit)
                                        if first_success_para is None:
                                            # Get first paragraph for comment placement
                                            for run in cell_edit['cell_runs']:
                                                para = run.get('para_elem')
                                                if para is not None:
                                                    first_success_para = para
                                                    break

                                    # Handle results based on success/failure counts
                                    if success_count > 0:
                                        # At least one cell succeeded - add overall comment
                                        if first_success_para is not None:
                                            # Try range comment only for full multi-cell success.
                                            # Otherwise keep stable reference-only behavior.
                                            if success_count == len(multi_cells):
                                                range_ok, range_reason = self._try_add_multi_cell_range_comment(
                                                    success_cells,
                                                    item.violation_reason,
                                                    item_author
                                                )
                                                if not range_ok:
                                                    fallback_comment_text = (
                                                        f"{{WHY}}{item.violation_reason}  "
                                                        f"{{WHERE}}{item.violation_text}"
                                                    )
                                                    self._append_reference_only_comment(
                                                        first_success_para,
                                                        fallback_comment_text,
                                                        item_author,
                                                        fallback_reason=range_reason
                                                    )
                                            else:
                                                self._append_reference_only_comment(
                                                    first_success_para,
                                                    item.violation_reason,
                                                    item_author
                                                )

                                        # Add individual comments for failed cells
                                        for cell_edit, cell_para, error_reason in failed_cells:
                                            if cell_para is not None:
                                                self._apply_cell_fallback_comment(
                                                    cell_para,
                                                    cell_edit['cell_violation'],
                                                    cell_edit['cell_revised'],
                                                    error_reason,
                                                    item
                                                )

                                        current_status = 'success'
                                        if self.verbose:
                                            if failed_cells:
                                                print(f"  [Multi-cell] Partial success: {success_count}/{len(multi_cells)} cells processed, {len(failed_cells)} failed")
                                            else:
                                                print(f"  [Multi-cell] Successfully applied track changes to all {len(multi_cells)} cells")
                                    else:
                                        # All cells failed - fallback to overall comment
                                        current_status = self._set_status_reason(
                                            'cross_cell_fallback',
                                            'CC_ALL_FAIL',
                                            'all extracted cells failed',
                                        )
                                else:
                                    # Changes cross cell boundaries - fallback to comment
                                    if self.verbose:
                                        print(f"  [Cross-cell] Changes cross cell boundaries, fallback to comment")
                                    current_status = self._set_status_reason(
                                        'cross_cell_fallback',
                                        'CC_XTRACT_FAIL',
                                        'cannot split cross-cell edits',
                                    )
                            if current_status is None and is_cross_row:
                                # Cell extraction failed for cross-row content
                                if self.verbose:
                                    print(f"  [Cross-row] Cell extraction failed, fallback to comment")
                                current_status = self._set_status_reason(
                                    'cross_row_fallback',
                                    'CR_XTRACT_FAIL',
                                    'cannot decompose cross-row edit',
                                )
                        elif len(para_elems) > 1:
                            # Actually spans multiple paragraphs
                            if self._is_table_mode(real_runs) or any(r.get('cell_elem') is not None for r in real_runs):
                                if self.verbose:
                                    print(f"  [Cross-paragraph] replace spans {len(para_elems)} paragraphs, fallback to comment")
                                # Debug: log snippet from marker on cross-paragraph fallback
                                try:
                                    if DEBUG_MARKER:
                                        combined_text = ''.join(r.get('text', '') for r in matched_runs_info)
                                        marker_idx = combined_text.find(DEBUG_MARKER)
                                        if marker_idx != -1:
                                            snippet_len = len(DEBUG_MARKER) + 60
                                            snippet = combined_text[marker_idx:marker_idx + snippet_len]
                                            print(f"  [DEBUG] Cross-paragraph content from marker: {repr(snippet)}")
                                except Exception:
                                    # Debug logging should never break apply flow
                                    pass
                                current_status = self._set_status_reason(
                                    'cross_paragraph_fallback',
                                    'CP_TBL_SPAN',
                                    f'replace spans {len(para_elems)} paras in table',
                                )
                            else:
                                if self.verbose:
                                    print(f"  [Cross-paragraph] replace spans {len(para_elems)} paragraphs, applying cross-paragraph replace")
                                current_status = self._apply_replace_cross_paragraph(
                                    matched_runs_info, current_start, violation_text,
                                    revised_text, item.violation_reason, item_author
                                )
                        else:
                            # All content is in single paragraph - safe to replace
                            if real_runs:
                                current_para = real_runs[0].get('para_elem', current_para)
                            current_status = self._apply_replace(
                                current_para, violation_text, revised_text,
                                item.violation_reason,
                                matched_runs_info, current_start,
                                item_author
                            )
                    else:
                        current_para = _resolve_target_para_for_match(
                            matched_runs_info, current_start, current_para
                        )
                        current_status = self._apply_replace(
                            current_para, violation_text, revised_text,
                            item.violation_reason,
                            matched_runs_info, current_start,
                            item_author
                        )

                return current_status, current_para

            success_status = None
            last_conflict_text = violation_text

            if item.fix_action in ('delete', 'replace'):
                success_status, target_para = _apply_delete_or_replace(
                    item.fix_action, matched_start, target_para
                )

                # Conflict recovery: try the next match in remaining block content.
                if success_status == 'conflict' and violation_text:
                    retry_runs_info = matched_runs_info
                    retry_text = ''.join(r.get('text', '') for r in retry_runs_info)

                    # Prefer full block range for retry search if available.
                    block_runs, block_text, _, block_boundary_error = self._collect_runs_info_across_paragraphs(
                        anchor_para, item.uuid_end
                    )
                    if block_runs and block_boundary_error is None:
                        retry_runs_info = block_runs
                        retry_text = block_text

                    search_from = matched_start
                    while True:
                        next_start = retry_text.find(violation_text, search_from + 1)
                        if next_start == -1:
                            break

                        matched_runs_info = retry_runs_info
                        matched_start = next_start
                        search_from = next_start
                        last_conflict_text = retry_text[next_start:next_start + len(violation_text)] or violation_text
                        target_para = _resolve_target_para_for_match(
                            matched_runs_info, matched_start, target_para
                        )

                        if self.verbose:
                            print(f"  [Conflict] Retrying next match at offset {matched_start}")

                        success_status, target_para = _apply_delete_or_replace(
                            item.fix_action, matched_start, target_para
                        )
                        if success_status != 'conflict':
                            break
            elif item.fix_action == 'manual':
                success_status = self._apply_manual(
                    target_para, violation_text,
                    item.violation_reason, revised_text,
                    matched_runs_info, matched_start,
                    item_author,
                    is_cross_paragraph  # Pass cross-paragraph flag
                )
            else:
                # Insert error comment for unknown action type
                self._apply_error_comment(anchor_para, item)
                return EditResult(False, item, f"Unknown action type: {item.fix_action}")

            status_reason = self._consume_status_reason(success_status)
            return self._build_result_from_status(
                item=item,
                success_status=success_status,
                target_para=target_para,
                anchor_para=anchor_para,
                violation_text=violation_text,
                revised_text=revised_text,
                matched_runs_info=matched_runs_info,
                matched_start=matched_start,
                item_author=item_author,
                is_cross_paragraph=is_cross_paragraph,
                numbering_stripped=numbering_stripped,
                last_conflict_text=last_conflict_text,
                status_reason=status_reason,
            )
                
        except Exception as e:
            # Insert error comment on exception if anchor paragraph exists
            if anchor_para is not None:
                self._apply_error_comment(anchor_para, item)
            return EditResult(False, item, str(e))
